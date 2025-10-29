#!/usr/bin/env python3
"""
Script de conversion automatique des fiches HAS
Markdown (.md) → Word (.docx)

EHPAD Belle Viste - Certification HAS novembre 2024
"""

import os
import sys

print("🔄 Conversion des fiches HAS : Markdown → Word")
print("=" * 60)

# Vérifier et installer Pandoc si nécessaire
try:
    import pypandoc

    # Essayer de télécharger Pandoc automatiquement
    try:
        pypandoc.get_pandoc_version()
        print("✓ Pandoc est déjà installé")
    except:
        print("📥 Téléchargement de Pandoc...")
        try:
            pypandoc.download_pandoc()
            print("✓ Pandoc téléchargé avec succès")
        except Exception as e:
            print(f"❌ Impossible de télécharger Pandoc : {e}")
            print("\n⚠️  Solution alternative : conversion basique")

            # Fallback : conversion simple sans Pandoc
            from docx import Document
            from docx.shared import Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            import re

            def convert_md_to_docx_simple(md_file, docx_file):
                """Conversion Markdown → Word simplifiée"""
                doc = Document()

                with open(md_file, 'r', encoding='utf-8') as f:
                    content = f.read()

                # Parser simple du Markdown
                lines = content.split('\n')

                for line in lines:
                    line = line.strip()

                    if not line:
                        continue

                    # Titre H1
                    if line.startswith('# '):
                        p = doc.add_heading(line[2:], level=1)
                        run = p.runs[0]
                        run.font.name = 'Arial'
                        run.font.size = Pt(18)
                        run.font.color.rgb = RGBColor(46, 83, 57)  # Vert Sapin

                    # Titre H2
                    elif line.startswith('## '):
                        p = doc.add_heading(line[3:], level=2)
                        run = p.runs[0]
                        run.font.name = 'Arial'
                        run.font.size = Pt(14)
                        run.font.color.rgb = RGBColor(139, 157, 131)  # Vert Sauge

                    # Titre H3
                    elif line.startswith('### '):
                        p = doc.add_heading(line[4:], level=3)

                    # Liste à puces
                    elif line.startswith('- ') or line.startswith('* '):
                        p = doc.add_paragraph(line[2:], style='List Bullet')

                    # Liste numérotée
                    elif re.match(r'^\d+\. ', line):
                        text = re.sub(r'^\d+\. ', '', line)
                        p = doc.add_paragraph(text, style='List Number')

                    # Ligne horizontale
                    elif line.startswith('---'):
                        doc.add_paragraph('_' * 60)

                    # Texte normal
                    else:
                        p = doc.add_paragraph(line)
                        run = p.runs[0] if p.runs else None
                        if run:
                            run.font.name = 'Arial'

                doc.save(docx_file)
                return True

            # Conversion avec méthode simple
            source_dir = "docs/has-certification/fiches-imperatifs"
            output_dir = os.path.join(source_dir, "word-output")
            os.makedirs(output_dir, exist_ok=True)

            md_files = [f for f in os.listdir(source_dir) if f.endswith('.md')]

            if not md_files:
                print(f"❌ Aucun fichier .md trouvé dans {source_dir}")
                sys.exit(1)

            print(f"\n📁 Dossier source : {source_dir}")
            print(f"📂 Dossier de sortie : {output_dir}")
            print(f"\n🔄 Conversion de {len(md_files)} fichiers...\n")

            for md_file in sorted(md_files):
                md_path = os.path.join(source_dir, md_file)
                docx_file = md_file.replace('.md', '.docx')
                docx_path = os.path.join(output_dir, docx_file)

                try:
                    convert_md_to_docx_simple(md_path, docx_path)
                    print(f"  ✓ {md_file} → {docx_file}")
                except Exception as e:
                    print(f"  ❌ Erreur avec {md_file}: {e}")

            print("\n" + "=" * 60)
            print("✅ Conversion terminée !")
            print(f"📊 {len(md_files)} fichiers convertis")
            print(f"📂 Fichiers Word dans : {output_dir}/")
            print("\n⚠️  Note : Conversion basique utilisée (tableaux non formatés)")
            print("💡 Conseil : Ouvrir les .docx dans Word pour ajuster le formatage")

            sys.exit(0)

except ImportError:
    print("❌ pypandoc n'est pas installé")
    print("Installation : pip3 install pypandoc")
    sys.exit(1)

# Si Pandoc fonctionne, utiliser pypandoc
print("\n🚀 Utilisation de Pandoc pour la conversion\n")

source_dir = "docs/has-certification/fiches-imperatifs"
output_dir = os.path.join(source_dir, "word-output")
os.makedirs(output_dir, exist_ok=True)

md_files = [f for f in os.listdir(source_dir) if f.endswith('.md')]

if not md_files:
    print(f"❌ Aucun fichier .md trouvé dans {source_dir}")
    sys.exit(1)

print(f"📁 Dossier source : {source_dir}")
print(f"📂 Dossier de sortie : {output_dir}")
print(f"\n🔄 Conversion de {len(md_files)} fichiers...\n")

count_success = 0
count_error = 0

for md_file in sorted(md_files):
    md_path = os.path.join(source_dir, md_file)
    docx_file = md_file.replace('.md', '.docx')
    docx_path = os.path.join(output_dir, docx_file)

    try:
        pypandoc.convert_file(
            md_path,
            'docx',
            outputfile=docx_path,
            extra_args=['--standalone']
        )
        print(f"  ✓ {md_file} → {docx_file}")
        count_success += 1
    except Exception as e:
        print(f"  ❌ Erreur avec {md_file}: {e}")
        count_error += 1

print("\n" + "=" * 60)
print("✅ Conversion terminée !")
print(f"📊 {count_success} fichiers convertis avec succès")
if count_error > 0:
    print(f"❌ {count_error} erreurs")
print(f"📂 Fichiers Word dans : {output_dir}/")
print("\n📝 Prochaines étapes :")
print("   1. Ouvrir les fichiers .docx dans Word")
print("   2. Appliquer la charte graphique Belle Viste")
print("   3. Personnaliser avec les noms des référents")
